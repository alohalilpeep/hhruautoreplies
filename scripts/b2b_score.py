"""Фильтр, обогащение и скоринг базы компаний для холодного аутрича DevOps-услуг.

    python scripts/b2b_score.py step1     предфильтр по названию и вакансиям
    python scripts/b2b_score.py step2     ИНН через DaData + численность/доходы из fns.db
    python scripts/b2b_score.py step2b    ИНН с сайтов для «ручных» (сайты — sites_out_*.json)
    python scripts/b2b_score.py step3     дофильтр, баллы, оффер, волны

Читает exports/b2b_2026-10-13.xlsx (лист «Все»), исходник не меняет.
Промежуточные результаты — exports/b2b_work/.
"""
import json
import random
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

EXPORTS = Path("/Users/mr_macbook/Documents/hh_autoreply/exports")
SRC = EXPORTS / "b2b_2026-10-13.xlsx"
WORK = EXPORTS / "b2b_work"


def low(s):
    return (s or "").lower().replace("ё", "е")


def rx(*words):
    return re.compile("|".join(words), re.I)


# ---------- списки исключений: проверяются в этом порядке ----------
# Посредники и вендоры — по точным брендам, раньше «гигантов»: РТК-ЦОД —
# вендор, хотя и «дочка» Ростелекома.
INTERMEDIARY = rx(
    r"nitka", r"\bibs\b", r"bell integrator", r"бэлл интегратор", r"usetech", r"ай-теко",
    r"andersen", r"neoflex", r"инфосистемы джет", r"jet infosystems", r"\bкрок\b", r"\bcroc\b",
    r"agima", r"hilbert", r"employcity", r"syberry", r"alta personnel", r"\bантал\b", r"antal",
    r"devhunt", r"easy ?hunt", r"getmatch", r"offer ?now", r"лига цифровой экономики",
    r"сателл", r"\bicl\b", r"ancor", r"анкор", r"itransition", r"it_one", r"it one\b",
    r"онланта", r"ланит", r"\bт1\b", r"\bt1\b", r"технократия", r"\bkts\b", r"emphasoft",
    r"live typing", r"кадров", r"рекрут", r"recruit", r"headhunt", r"\bhr\b", r"агентств",
    r"\bstaff", r"персонал", r"аутстаф", r"outstaff", r"аутсорс", r"outsourc", r"подбор",
    r"talent", r"hunt")
# неочевидные посредники: интеграторы и консалтинг без явного слова в названии
MAYBE_INTERMEDIARY = rx(r"интеграц", r"integrat", r"консалт", r"consult", r"solutions?\b",
                        r"софт ?лаб", r"soft ?lab", r"digital agency", r"веб-студ", r"студия")
VENDOR = rx(
    r"cloud\.ru", r"timeweb", r"таймвэб", r"т1 облако", r"\bастра\b", r"astra linux",
    r"группа компаний астра", r"orion soft", r"yadro", r"ядро", r"edge ?центр", r"edgecenter",
    r"ртк-цод", r"ddos-guard", r"stormwall", r"servicepipe", r"ngenix", r"servers\.com",
    r"ispsystem", r"\bсолар\b", r"solar\b", r"bi\.zone", r"positive technologies",
    r"позитив текнолоджиз", r"касперск", r"kaspersky", r"swordfish", r"group-ib", r"гарда",
    r"r-vision", r"р-вижн", r"infowatch", r"инфовотч", r"аренадата", r"arenadata", r"криптонит",
    r"qtech", r"\bбулат\b", r"selectel", r"селектел", r"\breg\.ru", r"рег\.ру", r"beget",
    r"mws\b", r"mts web services", r"yandex cloud", r"vk cloud", r"sbercloud", r"red ?os",
    r"ред ?софт", r"базальт", r"\bsber\b ?cloud", r"dataline", r"даталайн", r"ixcellerate",
    r"3data", r"linxdatacenter", r"\bкод безопасности\b", r"ufo ?team", r"ростелеком-солар")
GIANT = rx(
    # банки топ-20
    r"\bсбер", r"\bsber", r"\bвтб\b", r"\bvtb\b", r"альфа-банк", r"альфа банк", r"alfa-?bank",
    r"газпромбанк", r"россельхозбанк", r"\bрсхб", r"т-банк", r"тинькофф", r"tinkoff",
    r"московский кредитный", r"\bмкб\b", r"совкомбанк", r"открытие", r"райффайзен",
    r"raiffeisen", r"росбанк", r"промсвязьбанк", r"\bпсб\b", r"уралсиб", r"почта банк",
    r"ак барс", r"юникредит", r"unicredit", r"дом\.рф", r"мтс банк", r"хоум кредит",
    r"home credit", r"ренессанс банк", r"банк санкт-петербург", r"\bотп банк", r"otp bank",
    # телеком и «дочки»
    r"\bмтс\b", r"\bmts\b", r"билайн", r"beeline", r"вымпелком", r"\bt2\b", r"\bтеле2",
    r"tele2", r"ростелеком", r"rostelecom", r"\bртк\b", r"мегафон", r"megafon", r"ртлабс",
    # бигтех и ритейл
    r"яндекс", r"yandex", r"\bvk\b", r"вконтакте", r"\bozon", r"озон", r"wildberries",
    r"\brwb\b", r"\bwb\b", r"вайлдберриз", r"\bx5\b", r"магнит", r"magnit", r"лента",
    r"пятерочк", r"перекрест", r"авито", r"avito", r"kaspi",
    # госкорпорации и холдинги
    r"ростех", r"росатом", r"роскосмос", r"\bржд\b", r"аэрофлот", r"почта россии",
    r"норникель", r"норильск", r"русал", r"северсталь", r"евраз", r"газпром", r"лукойл",
    r"роснефть", r"сибур", r"транснефть", r"нлмк", r"\bммк\b", r"металлоинвест", r"фосагро",
    r"уралкалий", r"сургутнефтегаз", r"татнефть", r"\bполюс\b", r"мечел", r"\bсуэк\b",
    r"интер рао", r"россети", r"русгидро", r"\bафк система", r"\bсистема\b",
    # фастфуд
    r"вкусно.{0,3}и точка", r"\bkfc\b", r"rostic", r"burger king", r"бургер кинг")
GOV = rx(
    r"^\s*(гку|гуп|фгуп|гау|фку|фбун|гбу|фгбу|фгау|мку|мбу|гкоу|фгбоу)\b",
    r"\b(гку|гуп|фгуп|гау|фку|фбун|фгбу|фгау)\b", r"ано .*федеральн", r"федеральн\w* .*ано",
    r"правительств", r"министерств", r"департамент", r"\bнии\b", r"научно-исследовательск",
    r"\bнпо\b", r"концерн", r"оборон", r"метрополитен", r"росгвард", r"налогов\w* служб",
    r"федеральн\w* (служб|агентств|центр)", r"инфогород", r"мосгортранс", r"мфц",
    r"гнивц", r"фонд пенсионного", r"социального страхования", r"\bфнс\b", r"\bфсин\b",
    r"гознак")
LEGAL = re.compile(r"^\s*ип\s+|^индивидуальный предприниматель", re.I)
FIO = re.compile(r"^[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:вич|вна|ична|ич)$")
# «Не DevOps» — только если все вакансии про поддержку пользователей, телефонию
# или 1С. Системных и сетевых администраторов оставляем (решение 13.10.2026).
NOT_DEVOPS = rx(r"поддержк", r"support", r"helpdesk", r"хелпдеск", r"телефон", r"\bip-?телефон",
                r"\bаты\b", r"\b1с\b", r"\b1c\b", r"эникей", r"оператор", r"диспетчер")


# Слова, при которых вакансия остаётся, даже если в ней 1С или поддержка:
# «DevOps-инженер (1С)», «SRE (1С)», «Администратор 1С / Linux» — это наш клиент.
KEEP_WORDS = rx(r"devops", r"\bsre\b", r"reliability", r"platform", r"платформ", r"kubernetes",
                r"\bk8s\b", r"devsecops", r"mlops", r"ci/cd", r"инфраструктур", r"cloud", r"облач",
                r"linux", r"сервер", r"сетев", r"network", r"виртуал")


def all_not_devops(titles):
    return bool(titles) and all(NOT_DEVOPS.search(t) and not KEEP_WORDS.search(t)
                                for t in titles)


LOOKALIKE = str.maketrans("асеорхуктмнв", "aceopxyktmhb")


def unmix(s):
    """«Сloud.ru» с кириллической С → «cloud.ru»: в словах, где есть латиница,
    кириллические двойники латинских букв меняем на латиницу."""
    return re.sub(r"\S+", lambda m: m.group(0).translate(LOOKALIKE)
                  if re.search(r"[a-z]", m.group(0)) else m.group(0), s)


def classify(name, titles):
    n = low(name)
    variants = (n, unmix(n))
    if LEGAL.search(name) or FIO.match(name.strip()):
        return "посредник"
    if any(INTERMEDIARY.search(v) for v in variants):
        return "посредник"
    if any(VENDOR.search(v) for v in variants):
        return "инфра-вендор"
    if any(GIANT.search(v) for v in variants):
        return "гигант"
    if GOV.search(n):
        return "гос"
    if all_not_devops(titles):
        return "не DevOps"
    if MAYBE_INTERMEDIARY.search(n):
        return "посредник?"
    return ""


# ---------- склейка дублей ----------
ALIASES = {"мкб": "московскийкредитныйбанк"}
DROP_WORDS = re.compile(
    r"\b(ооо|оао|зао|пао|ао|нао|ип|llc|inc|ltd|limited|gmbh|corp|corporation|group|груп+а?|"
    r"гк|georgia|грузия|rnd|r&d|ру|ru|com|io|tech|техн?|team|тим)\b", re.I)


def dedupe_key(name):
    k = low(name)
    k = re.sub(r"[«»\"'`.,:;()\-–—/&]+", " ", k)
    k = DROP_WORDS.sub(" ", k)
    k = re.sub(r"\s+", "", k)
    return ALIASES.get(k, k)


def read_src():
    wb = load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Все"]
    rows = list(ws.iter_rows(values_only=True))
    head = [h for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        if d.get("Компания"):
            out.append(d)
    return out


def merge(rows):
    groups = {}
    for d in rows:
        groups.setdefault(dedupe_key(d["Компания"]), []).append(d)
    merged = []
    for key, items in groups.items():
        items.sort(key=lambda d: -(d.get("Вакансий инфраструктуры") or 0))
        base = dict(items[0])
        titles = []
        for d in items:
            for t in (d.get("Вакансии (до 5)") or "").split("\n"):
                if t.strip() and t.strip() not in titles:
                    titles.append(t.strip())
        stack = []
        for d in items:
            for s in (d.get("Стек") or "").split(","):
                if s.strip() and s.strip() not in stack:
                    stack.append(s.strip())
        dates_last = [d["Последняя инфра-вакансия"] for d in items if d.get("Последняя инфра-вакансия")]
        dates_old = [d["Самая старая открытая"] for d in items if d.get("Самая старая открытая")]
        sal_from = [d["Зарплата от"] for d in items if d.get("Зарплата от")]
        sal_to = [d["Зарплата до"] for d in items if d.get("Зарплата до")]
        base.update({
            "Вакансий инфраструктуры": sum(d.get("Вакансий инфраструктуры") or 0 for d in items),
            "Вакансии (до 5)": "\n".join(titles[:5]),
            "_все названия": titles,
            "Стек": ", ".join(stack[:10]),
            "Последняя инфра-вакансия": max(dates_last) if dates_last else None,
            "Самая старая открытая": min(dates_old) if dates_old else None,
            "Зарплата от": min(sal_from) if sal_from else None,
            "Зарплата до": max(sal_to) if sal_to else None,
            "Дубль с": ", ".join(d["Компания"] for d in items[1:]),
        })
        merged.append(base)
    return merged


def step1():
    rows = merge(read_src())
    for d in rows:
        d["Исключение"] = classify(d["Компания"], d["_все названия"])
    WORK.mkdir(exist_ok=True)
    out = WORK / "step1.json"
    out.write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    from collections import Counter
    stat = Counter(d["Исключение"] or "оставлено" for d in rows)
    print(f"строк в «Все»: {len(read_src())}, после склейки дублей: {len(rows)}")
    for k, v in stat.most_common():
        print(f"  {k:14s} {v}")
    dups = [d for d in rows if d["Дубль с"]]
    print(f"\nсклеено дублей: {len(dups)}")
    for d in dups:
        print(f"  {d['Компания']}  ⟵  {d['Дубль с']}")
    excl = [d for d in rows if d["Исключение"] and d["Исключение"] != "посредник?"]
    random.seed(13)
    print("\n30 случайных исключённых:")
    for d in random.sample(excl, min(30, len(excl))):
        t = (d["Вакансии (до 5)"] or "").split("\n")[0][:45]
        print(f"  [{d['Исключение']}] {d['Компания'][:38]:38s} | {t}")
    maybe = [d for d in rows if d["Исключение"] == "посредник?"]
    print(f"\nвсе «посредник?» ({len(maybe)}):")
    for d in maybe:
        t = (d["Вакансии (до 5)"] or "").split("\n")[0][:45]
        print(f"  {d['Компания'][:40]:40s} | {t}")
    print(f"\nрезультат шага 1: {out}")


# ---------- шаг 2: обогащение через DaData ----------
DADATA = "https://suggestions.dadata.ru/suggestions/api/4_1/rs"
CACHE = WORK / "dadata_cache.json"
OPF_RE = re.compile(r"^(ооо|оао|зао|пао|ао|нао|ип|ткп|llc|ltd|inc)\s+|[«»\"']", re.I)


def dadata_key():
    env = Path("/Users/mr_macbook/Documents/hh_autoreply/.env").read_text()
    m = re.search(r"^DADATA_API_KEY=(\S+)", env, re.M)
    if not m:
        raise SystemExit("нет DADATA_API_KEY в .env первого аккаунта")
    return m.group(1)


def cached(cache, kind, query, call):
    k = f"{kind}:{query}"
    if k not in cache:
        cache[k] = call()
        CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
        import time
        time.sleep(0.15)                       # лимит DaData — 20 запросов в секунду
    return cache[k]


FNS_DB = WORK / "fns.db"
DEAD = {"LIQUIDATED", "LIQUIDATING", "BANKRUPT"}
# «Гознак. Информационные технологии», «СОГАЗ: г. Москва», «T2. IT и Digital», «… | IT»
DIVISION_SUFFIX = re.compile(
    r"(\.\s+[A-ZА-ЯЁ].*|:\s.*|\s+\|.*|,\s*(г\.|компания|сеть|работа).*|"
    r"\s+(направление|it|ит|it & digital|digital|r&d)\b.*)$", re.I)


TO_CYR = str.maketrans("aceopxyktmhb", "асеорхуктмнв")


def name_tokens(s):
    """Слова названия для сравнения; латинские двойники букв → кириллица,
    чтобы «1C-Рарус» совпал с «1С-РАРУС», а Okko — с «ОККО»."""
    s = OPF_RE.sub(" ", low(s)).translate(TO_CYR)
    return {t for t in re.split(r"[^\w]+", s) if len(t) > 1}


def queries(company):
    """Запросы к DaData по очереди: основной без суффикса подразделения, затем запасные."""
    base = DIVISION_SUFFIX.sub("", company).strip()
    out = [re.sub(r"\s*\(.*?\)", "", base).strip()]
    for p in re.findall(r"\((.*?)\)", company):          # «FunFlow (ООО Мобильные Игровые Решения)»
        out.append(p)
    out.append(re.sub(r"\.(ru|рф|com|io|tech)\b", "", out[0], flags=re.I))   # «Travelata.ru»
    out.append(" ".join(re.findall(r"[A-Za-z][A-Za-z0-9&-]+", out[0])))       # латинский бренд
    seen, res = set(), []
    for q in out:
        q = OPF_RE.sub(" ", q).strip(" .,-")
        if len(q) >= 2 and low(q) not in seen:
            seen.add(low(q))
            res.append(q)
    return res


def fns_row(db, inn):
    r = db.execute("SELECT staff, income FROM fns WHERE inn=?", (inn,)).fetchone()
    return r or (None, None)


def is_it(okved):
    return bool(re.match(r"6[23]\.", okved or ""))


def match(query, suggestions, db):
    """Выбор юрлица (правила 13.10.2026). Слова бренда — из очищенного запроса.
    «высокая» — ровно один действующий кандидат, в названии которого есть все слова
    бренда; численность не проверяем (< 10 — пометка, нет в fns.db — пустые данные).
    Иначе гибрид, только среди точных совпадений: численность ≥ 10 по ФНС; один с
    ОКВЭД 62/63 → «авто»; несколько → самый крупный, если он в 5 раз больше каждого
    из остальных. Кандидат без данных ФНС делает выбор неоднозначным. Регион не учитываем.
    Возвращает (лучший или None, уверенность, топ-3, причина)."""
    brand = name_tokens(query)
    pool, seen = [], set()
    for s in suggestions:
        d = s["data"]
        inn = d.get("inn")
        if not inn or inn in seen or d.get("type") == "INDIVIDUAL":
            continue
        seen.add(inn)
        if (d.get("state") or {}).get("status") in DEAD:
            continue
        cand = name_tokens(s["value"]) | name_tokens(((d.get("name") or {}).get("short")) or "")
        if not brand & cand:                              # подсказка DaData «по созвучию»
            continue
        staff, income = fns_row(db, inn)
        pool.append({"s": s, "staff": staff, "income": income, "exact": cand >= brand})
    pool.sort(key=lambda c: -(c["staff"] or 0))
    if not pool:
        return None, "ручная", [], "нет кандидатов"
    exact = [c for c in pool if c["exact"]]
    if len(exact) == 1:
        return exact[0], "высокая", pool[:3], ""
    if not exact:
        return None, "ручная", pool[:3], "нет точного совпадения названия"
    unknown = [c for c in exact if c["staff"] is None]
    sized = [c for c in exact if c["staff"] is not None and c["staff"] >= 10]
    it = [c for c in sized if is_it(c["s"]["data"].get("okved"))]
    it_unknown = [c for c in unknown if is_it(c["s"]["data"].get("okved"))]
    if len(it) == 1 and not it_unknown:
        return it[0], "авто", exact[:3], ""
    if sized and not unknown and sized[0]["staff"] >= 5 * max(c["staff"] for c in exact[1:]):
        return sized[0], "авто", exact[:3], ""
    return None, "ручная", exact[:3], "несколько однофамильцев, ни один не выделяется"


def okved_name(cache, key, code):
    if not code:
        return ""
    r = cached(cache, "okved", code, lambda: requests_post(
        key, "findById/okved2", {"query": code}))
    s = (r or {}).get("suggestions") or []
    return s[0]["data"].get("name", "") if s else ""


def requests_post(key, method, body):
    import requests
    r = requests.post(f"{DADATA}/{method}", json=body, timeout=20,
                      headers={"Authorization": f"Token {key}", "Content-Type": "application/json"})
    return r.json() if r.ok else {"error": r.status_code}


def region(d):
    return ((d.get("address") or {}).get("data") or {}).get("region_with_type") or ""


def describe(c):
    d = c["s"]["data"]
    return (f"{d.get('inn')} | {c['s']['value']} | {d.get('okved') or '—'} | "
            f"{c['staff']} чел. | {region(d) or '—'}")


def step2(limit=None):
    import sqlite3
    from collections import Counter
    rows = json.loads((WORK / "step1.json").read_text(encoding="utf-8"))
    key = dadata_key()
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    db = sqlite3.connect(FNS_DB)
    todo = [r for r in rows if r["Исключение"] in ("", "посредник?")]
    if limit:
        todo = todo[:limit]
    for i, r in enumerate(todo, 1):
        tried = []
        for q in queries(r["Компания"]):
            res = cached(cache, "party10", q, lambda: requests_post(
                key, "suggest/party", {"query": q, "count": 10}))
            tried.append(q)
            best, conf, top, why = match(q, (res or {}).get("suggestions") or [], db)
            if best or top:                 # нашли кандидатов — запасные запросы не нужны
                break
        r.update({"ИНН": "", "Юрлицо": "", "ОКВЭД код": "", "ОКВЭД расшифровка": "",
                  "Регион": "", "Статус": "", "Численность": None, "Доходы (ФНС)": None,
                  "Пометка": "", "Сайт": r.get("Сайт") or "",
                  "Уверенность сопоставления": conf, "Причина ручной проверки": why,
                  "Кандидаты (топ-3)": "\n".join(describe(c) for c in top) if not best else "",
                  "Запросы DaData": " / ".join(tried)})
        if best:
            d = best["s"]["data"]
            code = d.get("okved") or ""
            r.update({
                "ИНН": d.get("inn"), "Юрлицо": best["s"]["value"], "ОКВЭД код": code,
                "ОКВЭД расшифровка": okved_name(cache, key, code), "Регион": region(d),
                "Статус": "действующая", "Численность": best["staff"],
                "Доходы (ФНС)": best["income"],
                "Пометка": "нет данных ФНС по юрлицу" if best["staff"] is None else
                           "мало сотрудников в юрлице, возможно, штат в другом юрлице группы"
                           if best["staff"] < 10 else ""})
        if i % 50 == 0:
            print(f"  … {i}/{len(todo)}", flush=True)
    out = WORK / ("step2_sample.json" if limit else "step2.json")
    out.write_text(json.dumps(rows if not limit else todo, ensure_ascii=False, indent=1,
                              default=str), encoding="utf-8")

    stat = Counter(r["Уверенность сопоставления"] for r in todo)
    print(f"\nобогащено строк: {len(todo)}")
    for k in ("высокая", "авто", "ручная"):
        print(f"  {k:8s} {stat.get(k, 0)}")
    why = Counter(r["Причина ручной проверки"] for r in todo if r["Причина ручной проверки"])
    print("причины «ручная»:")
    for k, v in why.most_common():
        print(f"  {v:4d}  {k}")
    auto = [r for r in todo if r["Уверенность сопоставления"] == "авто"]
    random.seed(13)
    print("\n15 случайных «авто»:")
    for r in random.sample(auto, min(15, len(auto))):
        inc = f"{r['Доходы (ФНС)'] / 1e6:,.0f} млн".replace(",", " ") if r["Доходы (ФНС)"] else "—"
        print(f"  {r['Компания'][:32]:32s} → {r['Юрлицо'][:40]:40s} ИНН {r['ИНН']} | "
              f"{r['ОКВЭД код']:8s} | {r['Численность']:>5} чел. | {inc:>10s} | {r['Регион'][:20]}")
    print(f"\nрезультат: {out}")


# ---------- шаг 2.1: ИНН с сайта ----------
# Сайты нашёл веб-поиск (sites_out_*.json): {company, site, inn, legal, foreign, note}.
SITE_CACHE = WORK / "site_cache.json"
SITE_PAGES = ("", "/contacts", "/requisites", "/about", "/privacy", "/policy", "/oferta")
NOT_RF = "не РФ / нет юрлица РФ"
INN_RE = re.compile(r"(?:ИНН|INN)[\s:№/]*(?:КПП[\s:№]*\d{9}[\s,;/]*)?(\d{12}|\d{10})(?!\d)")
OGRN_RE = re.compile(r"(?:ОГРН(?:ИП)?|OGRN)[\s:№]*(\d{15}|\d{13})(?!\d)")
# «не РФ» — только по форме юрлица (правило 13.10.2026), страны в тексте не считаются:
# на российском сайте может быть «офис в Казахстане».
FOREIGN_FORM = r"(?<![\w.])(ТОО|TOO|LLC|L\.L\.C\.?|Ltd\.?|LTD|GmbH|ЧК|Inc\.?|B\.V\.)(?![\w])"
FOREIGN_RE = re.compile(FOREIGN_FORM)
CAPTCHA_RE = re.compile(r"captcha|cf-chl|challenge-platform|ddos-guard|are you a robot|"
                        r"подтвердите, что вы не робот", re.I)


def inn_ok(inn):
    """Контрольная сумма ИНН — отсекает случайные 10/12 цифр."""
    d = [int(c) for c in inn]
    def cs(w):
        return sum(a * b for a, b in zip(w, d)) % 11 % 10
    if len(d) == 10:
        return cs([2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[9]
    if len(d) == 12:
        return (cs([7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[10]
                and cs([3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[11])
    return False


import threading
SITE_LOCK = threading.Lock()
_site_writes = [0]


def save_site_cache(cache):
    with SITE_LOCK:
        SITE_CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")


def fetch_site(cache, url):
    """Текст страницы (или маркер ошибки), с кэшем и паузой 1.5 с. Потокобезопасно:
    кэш на диск — раз в 20 новых страниц и в конце шага."""
    import time
    import requests
    with SITE_LOCK:
        if url in cache:
            return cache[url]
    try:
        r = requests.get(url, timeout=(5, 15), headers={
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/128.0 Safari/537.36",
            "Accept-Language": "ru-RU,ru;q=0.9"})
        text = r.text if r.ok else f"#HTTP {r.status_code}"
        if CAPTCHA_RE.search(r.text[:20000]) and len(r.text) < 60000:
            text = "#CAPTCHA"
    except Exception as e:                         # noqa: BLE001 — сайт недоступен
        text = f"#ERR {type(e).__name__}"
    text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;|&#160;|\s+", " ", text)
    text = text[-60000:]                           # реквизиты почти всегда в футере
    with SITE_LOCK:
        cache[url] = text
        _site_writes[0] += 1
        if _site_writes[0] % 20 == 0:
            SITE_CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    time.sleep(1.5)
    return text


def site_inn(cache, site):
    """(ИНН или ОГРН, признак иностранного юрлица, капча?) по главной и служебным страницам."""
    from collections import Counter
    base = re.match(r"https?://[^/]+", site.strip()).group(0)
    found, foreign = Counter(), False
    for i, p in enumerate(SITE_PAGES):
        text = fetch_site(cache, base + p)
        if text == "#CAPTCHA":
            return None, False, True
        if i == 0 and text.startswith("#"):        # главная не открылась — дальше не идём
            break
        for inn in INN_RE.findall(text):
            if inn_ok(inn):
                found[inn] += 1
        for ogrn in OGRN_RE.findall(text):
            found["ogrn:" + ogrn] += 1
        if FOREIGN_RE.search(text):
            foreign = True
        inns = [k for k in found if not k.startswith("ogrn:")]
        if inns and i >= 1:                        # главная + ещё одна страница — хватит
            break
    inns = [k for k, _ in found.most_common() if not k.startswith("ogrn:")]
    if inns:
        return inns[0], foreign, False
    ogrns = [k[5:] for k, _ in found.most_common() if k.startswith("ogrn:")]
    return (ogrns[0] if ogrns else None), foreign, False


TRANSLIT = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
                    ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "y", "k", "l", "m", "n",
                     "o", "p", "r", "s", "t", "u", "f", "h", "ts", "ch", "sh", "sch", "", "y", "",
                     "e", "yu", "ya"]))


def guess_sites(company):
    """Домены-кандидаты, когда веб-поиск до компании не дошёл: бренд латиницей
    (как есть или транслитом) + .ru / .com / .io / .tech."""
    base = re.sub(r"\s*\(.*?\)", "", DIVISION_SUFFIX.sub("", company))
    base = OPF_RE.sub(" ", base).strip(" .,-")
    if re.search(r"\.(ru|com|io|tech|рф)\b", base, re.I):
        return ["https://" + re.sub(r"\s+", "", base.lower()) + "/"]
    words = [w for w in re.split(r"[^\w-]+", low(base)) if w]
    if not words or len(words) > 3:
        return []
    lat = "".join("".join(TRANSLIT.get(c, c) for c in w) for w in words)
    lat = re.sub(r"[^a-z0-9-]", "", lat)
    if len(lat) < 3:
        return []
    return [f"https://{lat}.{tld}/" for tld in ("ru", "com", "io", "tech")]


def brand_on_page(company, text):
    """Сайт «наш», если на главной есть хоть одно значимое слово бренда (≥ 4 букв)."""
    words = [w for w in name_tokens(DIVISION_SUFFIX.sub("", company)) if len(w) >= 4]
    page = low(text).translate(TO_CYR)
    return bool(words) and any(w in page for w in words)


def find_party(cache, key, ident):
    res = cached(cache, "byid", ident, lambda: requests_post(
        key, "findById/party", {"query": ident, "count": 1}))
    s = (res or {}).get("suggestions") or []
    return s[0] if s else None


def step2b():
    import sqlite3
    from collections import Counter
    rows = json.loads((WORK / "step2.json").read_text(encoding="utf-8"))
    sites = {}
    for f in sorted(WORK.glob("sites_out_*.json")):
        for x in json.loads(f.read_text(encoding="utf-8")):
            sites[x["company"]] = x
    key = dadata_key()
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    scache = json.loads(SITE_CACHE.read_text(encoding="utf-8")) if SITE_CACHE.exists() else {}
    db = sqlite3.connect(FNS_DB)
    stat = Counter()
    todo = [r for r in rows if r.get("Уверенность сопоставления") == "ручная"]

    def crawl(r):
        """Сетевая часть по одной компании: сайт (угаданный, если поиск не дошёл) и ИНН с него."""
        x = sites.get(r["Компания"]) or {"company": r["Компания"], "note": "лимит"}
        unsearched = not x.get("site") and re.search(r"лимит", x.get("note") or "")
        if unsearched:                             # поиск не дошёл — пробуем угадать домен
            for url in guess_sites(r["Компания"]):
                text = fetch_site(scache, url)
                if not text.startswith("#") and brand_on_page(r["Компания"], text):
                    x = dict(x, site=url, guessed=True)
                    break
        # флаг поиска «foreign» сам по себе не доказательство — нужна форма юрлица
        ident, captcha, src = None, False, ""
        foreign = bool(x.get("foreign") and FOREIGN_RE.search(x.get("legal") or ""))
        if x.get("inn") and inn_ok(str(x["inn"])):
            ident, src = str(x["inn"]), "поиск"
        elif x.get("site"):
            ident, f2, captcha = site_inn(scache, x["site"])
            foreign = foreign or (f2 and not ident)
            src = "сайт (домен угадан)" if x.get("guessed") else "сайт"
        return x, bool(unsearched), ident, foreign, captcha, src

    # сайты разные, поэтому качаем по 8 параллельно; пауза 1.5 с — внутри каждого потока
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(8) as pool:
        crawled = []
        for i, res in enumerate(pool.map(crawl, todo), 1):
            crawled.append(res)
            if i % 25 == 0:
                print(f"  … {i}/{len(todo)}", flush=True)
    save_site_cache(scache)

    for r, (x, unsearched, ident, foreign, captcha, src) in zip(todo, crawled):
        if unsearched and not x.get("site"):
            r["Причина ручной проверки"] = "сайт не искали (лимит поиска), домен не угадан"
            stat["не искали, домен не угадан"] += 1
            continue
        if unsearched:
            stat["домен угадан"] += 1
        if x.get("site"):
            r["Сайт"] = x["site"]
        if captcha:
            stat["капча на сайте"] += 1
            r["Пометка"] = "капча на сайте, ИНН не проверен"
            continue
        party = find_party(cache, key, ident) if ident else None
        if party and (party["data"].get("state") or {}).get("status") not in DEAD:
            d = party["data"]
            staff, income = fns_row(db, d.get("inn"))
            code = d.get("okved") or ""
            r.update({
                "ИНН": d.get("inn"), "Юрлицо": party["value"], "ОКВЭД код": code,
                "ОКВЭД расшифровка": okved_name(cache, key, code), "Регион": region(d),
                "Статус": "действующая", "Численность": staff, "Доходы (ФНС)": income,
                "Уверенность сопоставления": "сайт", "Причина ручной проверки": "",
                "Кандидаты (топ-3)": "", "Источник ИНН": src,
                "Пометка": "нет данных ФНС по юрлицу" if staff is None else
                           "мало сотрудников в юрлице, возможно, штат в другом юрлице группы"
                           if staff < 10 else ""})
            stat[f"ИНН найден ({src})"] += 1
        elif party:
            stat["ИНН с сайта — ликвидировано"] += 1
        elif foreign:
            r["Регион"] = NOT_RF
            r["Причина ручной проверки"] = "юрлицо не РФ"
            stat["не РФ (иностранное юрлицо подтверждено)"] += 1
        elif not x.get("site"):
            r["Причина ручной проверки"] = "сайт не найден"
            stat["сайт не найден"] += 1
        else:
            stat["сайт есть, ИНН на нём не найден"] += 1
    (WORK / "step2b.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str),
                                     encoding="utf-8")
    print(f"\nстрок на поиск по сайтам: {len(todo)}")
    for k, v in stat.most_common():
        print(f"  {v:4d}  {k}")
    print(f"результат: {WORK / 'step2b.json'}")


# ---------- шаги 3–4: скоринг и оффер ----------
TODAY = "2026-10-13"
LEAD_RE = rx(r"\bhead\b", r"\blead\b", r"руковод", r"тимлид", r"team ?lead", r"tech ?lead",
             r"архитект", r"architect", r"единственн", r"первый devops", r"first devops")
JUNIOR_RE = rx(r"junior", r"джун", r"младш", r"стаж", r"trainee", r"intern", r"ученик")
FOREIGN_NAME = re.compile(FOREIGN_FORM, re.I)
# DevOps-фильтр для волн 1–2 — по НАЗВАНИЯМ вакансий; «инфраструктур» само по себе не в счёт
DEVOPS_TITLE = rx(r"devops", r"\bsre\b", r"site reliability", r"platform engineer", r"платформ",
                  r"kubernetes", r"\bk8s\b", r"devsecops", r"mlops", r"ci/cd",
                  r"релиз.?инженер", r"release engineer", r"инфраструктуры и автоматизации")
# сисадмин с таким стеком — максимум волна 2
DEVOPS_STACK = rx(r"ansible", r"terraform", r"docker", r"kubernetes", r"\bk8s\b")
# Дочки крупных групп (DaData бесплатно учредителей не отдаёт): ИНН или начало
# названия компании → группа. Только те, в принадлежности которых уверен (13.10.2026).
GROUP_CHILD = {
    "7716971253": "Аэрофлот",                  # АФЛТ-Системс
    "7736249247": "Сбер",                      # Домклик
    "7804568396": "Сбер",                      # Салют для Бизнеса
    "7729405872": "Татнефть",                  # Банк Зенит
    "7708119944": "Роснефть",                  # ИК Сибинтек
    "ао центр биометрических": "Ростелеком",   # ЦБТ, оператор ЕБС
    "росэлторг": "ВТБ",
    "ао «силовые машины»": "Севергрупп",
    "пао т плюс": "Ренова",
    "ашан": "Auchan",
}


# Ручные решения пользователя по конкретным компаниям (14.10.2026), ключ — low(«Компания»).
MANUAL = {
    "честный знак.рф": {"Исключение": "гос"},
    "бкс it & digital": {"Исключение": "гигант"},
    "бюро 1440": {"Исключение": "гигант"},
    "rubytech": {"Исключение": "посредник"},
    "ооо отр": {"wave": 3, "note": "господрядчик"},
    "ооо эвотор": {"wave": 2, "size_unknown": True, "note": "размер не проверен"},
    "ооо клауд солюшенс": {"note": "проверить: облако/интегратор? сайт не найден"},
    # авто-сопоставление взяло региональную франшизу «2ГИС.ПЕНЗА», а не головное юрлицо
    "2гис": {"unmatch": True, "note": "ИНН снят: найдено юрлицо франшизы 2ГИС.ПЕНЗА"},
}


def group_of(r):
    if r.get("ИНН") in GROUP_CHILD:
        return GROUP_CHILD[r["ИНН"]]
    name = low(r["Компания"])
    return next((g for k, g in GROUP_CHILD.items() if not k.isdigit() and name.startswith(k)), "")
OFFERS = [
    (1, "Chaos-тест отказоустойчивости", rx(r"\bsre\b", r"reliability", r"high availability",
                                             r"\bha\b", r"инцидент", r"надежн", r"надёжн")),
    (2, "Мониторинг и алертинг за 3 дня", rx(r"мониторинг", r"observability", r"prometheus",
                                              r"grafana", r"zabbix", r"\belk\b", r"логи",
                                              r"victoriametrics", r"opensearch", r"loki")),
    # сетевую и инфраструктурную безопасность не берём: это не про CI/CD
    (3, "Безопасность в CI/CD", rx(r"devsecops", r"appsec", r"application security",
                                    r"безопасн\w* разработ", r"\bsast\b", r"\bdast\b",
                                    r"gitleaks", r"trivy", r"sonarqube")),
    (4, "FinOps-аудит облаков", rx(r"\baws\b", r"\bgcp\b", r"azure", r"yandex ?cloud",
                                    r"яндекс ?облак", r"оптимизац\w* затрат", r"finops")),
    (5, "Пилотная миграция 1 сервиса", rx(r"миграц", r"переезд", r"импортозамещ", r"astra",
                                          r"астра", r"red ?os", r"ред ?ос")),
    # «системный администратор» сам по себе не триггер (правило 13.10.2026)
    (6, "Ansible для железа", rx(r"ansible", r"bare.?metal", r"серверн\w* оборудован",
                                  r"виртуализ", r"vmware", r"proxmox")),
    (7, "Единое dev-окружение на Docker", rx(r"помощь разработчик", r"рост команды", r"онбординг",
                                             r"docker", r"dev.?окружен")),
    (8, "Тест-стенды на каждый PR", rx(r"тестов\w* сред", r"\bstage\b", r"staging", r"qa.?окружен",
                                        r"стенд")),
    (9, "Helm-шаблонизация деплоя", rx(r"kubernetes", r"\bk8s\b", r"helm", r"openshift",
                                        r"deckhouse")),
    (10, "Ускорение CI/CD", rx(r"ci/cd", r"\bci\b", r"gitlab", r"jenkins", r"teamcity",
                               r"релиз", r"release", r"1с devops", r"1c devops")),
]
MLOPS = rx(r"mlops", r"\bml\b")
WHO = ((100, "CTO / основатель"), (500, "CTO / техдир"),
       (1000, "руководитель инфраструктуры/платформы или CIO"))


def as_date(v):
    return str(v)[:10] if v else ""


def offer(titles, stack):
    """(главный, запасной, зацепка, есть ли явный триггер)."""
    hits = []
    for num, name, pat in OFFERS:
        t_hits = [t for t in titles if pat.search(t)]
        s_hits = [s.strip() for s in (stack or "").split(",") if pat.search(s)]
        score = 2 * len(t_hits) + len(s_hits)
        if score:
            hits.append((score, -num, num, name, t_hits[0] if t_hits else s_hits[0]))
    hits.sort(reverse=True)
    if not hits and any(MLOPS.search(t) for t in titles):
        return "9. Helm-шаблонизация деплоя", "10. Ускорение CI/CD", titles[0], True
    if not hits:
        return ("10. Ускорение CI/CD", "2. Мониторинг и алертинг за 3 дня",
                titles[0] if titles else "", False)
    main = f"{hits[0][2]}. {hits[0][3]}"
    spare = f"{hits[1][2]}. {hits[1][3]}" if len(hits) > 1 else (
        "10. Ускорение CI/CD" if hits[0][2] != 10 else "2. Мониторинг и алертинг за 3 дня")
    return main, spare, hits[0][4], True


def score(r):
    """Баллы шага 3. Размерные пункты при неизвестной численности — 0 (правило 13.10.2026).
    Возвращает (балл, балл без размера, «почему»)."""
    from datetime import date
    titles = r.get("_все названия") or []
    staff = r.get("_размер")
    pts, why, size_pts = 0, [], 0
    old = as_date(r.get("Самая старая открытая"))
    if old and (date.fromisoformat(TODAY) - date.fromisoformat(old)).days >= 21:
        pts += 3
        why.append("+3 висит 21+ дн")
    if any(LEAD_RE.search(t) for t in titles):
        pts += 3
        why.append("+3 лид/первый DevOps")
    if (r.get("Вакансий инфраструктуры") or 0) >= 2 and staff is not None and staff <= 1000:
        size_pts += 2
        why.append("+2 2+ вакансии при ≤1000")
    if staff is not None and 50 <= staff <= 1000:
        size_pts += 2
        why.append("+2 численность 50–1000")
    if (r.get("Зарплата до") or 0) >= 300000:
        pts += 1
        why.append("+1 вилка до ≥300к")
    if r.get("_триггер"):
        pts += 1
        why.append("+1 триггер под оффер")
    if titles and all(JUNIOR_RE.search(t) for t in titles):
        pts -= 3
        why.append("−3 только junior")
    if r.get("Регион") == NOT_RF or FOREIGN_NAME.search(r["Компания"]):
        pts -= 2
        why.append("−2 не РФ")
    if (r.get("ОКВЭД код") or "").startswith("64.19") and (staff or 0) > 1000:
        pts -= 2
        why.append("−2 банк >1000")
    return pts + size_pts, pts, "; ".join(why) or "0"


def post_filter(r):
    """Дофильтр после обогащения: ОКВЭД и численность."""
    code, staff = r.get("ОКВЭД код") or "", r.get("Численность")
    if r["Исключение"]:
        return r["Исключение"]
    if code.startswith("78"):
        return "посредник"
    if code.startswith("84"):
        return "гос"
    if staff is not None and staff > 3000:
        return "гигант"
    if group_of(r):
        return f"гигант (группа {group_of(r)})"
    if code.startswith("70.22"):
        return "посредник?"
    return ""


def step3():
    from collections import Counter
    src = WORK / "step2b.json"
    if not src.exists():                           # шаг 2.1 ещё не прогнан
        src = WORK / "step2.json"
    rows = json.loads(src.read_text(encoding="utf-8"))
    for r in rows:
        manual = MANUAL.get(low(r["Компания"]), {})
        if manual.get("unmatch"):
            r.update({"ИНН": "", "Юрлицо": "", "ОКВЭД код": "", "ОКВЭД расшифровка": "",
                      "Регион": "", "Численность": None, "Доходы (ФНС)": None,
                      "Уверенность сопоставления": "ручная"})
        r["Исключение"] = (manual["Исключение"] + " (решение вручную)" if "Исключение" in manual
                           else post_filter(r))
        if r["Исключение"] and r["Исключение"] != "посредник?":
            r["Волна"] = ""
            r["Лист"] = "Исключены"
            continue
        titles = r.get("_все названия") or []
        r["Оффер"], r["Запасной оффер"], r["Зацепка"], r["_триггер"] = offer(titles, r.get("Стек"))
        # < 10 человек в юрлице — скорее не то юрлицо группы: размер считаем неизвестным
        staff = r.get("Численность")
        if (staff is not None and staff < 10) or manual.get("size_unknown"):
            staff = None
        r["_размер"] = staff
        r["Балл"], r["Балл без размера"], r["Почему"] = score(r)
        r["Кого искать"] = "CTO" if staff is None else next(
            (w for lim, w in WHO if staff < lim), "руководитель инфраструктуры/платформы или CIO")
        notes = [r.get("Пометка")] if r.get("Пометка") else []
        devops = any(DEVOPS_TITLE.search(t) for t in titles)
        r["DevOps-фильтр"] = "да" if devops else "нет"
        if staff is None:
            # размер пользователь увидит сам, когда будет искать ЛПРа (правило 13.10.2026)
            if devops and r["Балл без размера"] >= 4:
                wave = 2
                notes.append("размер не проверен")
            else:
                wave = 3
                notes.append("размер не определён")
        elif staff < 50:
            wave = 3
            notes.append("микро/малая")
        else:
            wave = 1 if r["Балл"] >= 7 else 2 if r["Балл"] >= 4 else 3
        if not devops:
            if DEVOPS_STACK.search(r.get("Стек") or ""):
                wave = max(wave, 2)                # сисадмин с Ansible/Terraform/Docker/k8s
                notes.append("сисадмин, DevOps-стек")
            else:
                wave = 3
                notes.append("сисадмин")
        if "wave" in manual:
            wave = manual["wave"]
            notes = [n for n in notes if n not in ("размер не определён", "размер не проверен")]
        if manual.get("note"):
            notes.append(manual["note"])
        r["Волна"], r["Лист"] = wave, f"Волна {wave}"
        if r["Исключение"] == "посредник?":
            r["Лист"] = "Ручная проверка"
            notes.append("посредник?")
        r["Пометка"] = "; ".join(dict.fromkeys(notes))
    (WORK / "step3.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str),
                                    encoding="utf-8")

    sheets = Counter(r["Лист"] for r in rows)
    print("по листам:")
    for k in ("Волна 1", "Волна 2", "Волна 3", "Ручная проверка", "Исключены"):
        print(f"  {k:16s} {sheets.get(k, 0)}")
    conf = Counter(r.get("Уверенность сопоставления") for r in rows
                   if r.get("Уверенность сопоставления"))
    print("уверенность сопоставления:", dict(conf))
    w3 = Counter(n for r in rows if r["Лист"] == "Волна 3"
                 for n in ("сисадмин", "размер не определён", "микро/малая")
                 if n in r["Пометка"].split("; "))
    print("волна 3 по пометкам:", dict(w3))
    for wave in ("Волна 1", "Волна 2"):
        ws = sorted((r for r in rows if r["Лист"] == wave),
                    key=lambda r: (-r["Балл"], as_date(r.get("Самая старая открытая"))))
        print(f"\n{wave} ({len(ws)}):")
        for r in ws:
            print(f"  {r['Балл']:2d} | {r['Компания'][:28]:28s} | {r['Численность'] or '—':>5} | "
                  f"{r['Почему'][:70]:70s} | {r['Оффер'][:30]:30s} | {r['Зацепка'][:45]}"
                  + (f" | {r['Пометка']}" if r["Пометка"] else ""))
    print(f"результат: {WORK / 'step3.json'}")


# ---------- шаг 5: b2b_scored.xlsx ----------
OUT = EXPORTS / "b2b_scored.xlsx"
LPR_COLS = ["ЛПР имя", "ЛПР должность", "Telegram ЛПР", "LinkedIn/почта", "Контакт рекрутера",
            "Статус", "Дата касания", "Следующий шаг", "Комментарий"]
COLS = ["Компания", "Волна", "Волна вручную", "Балл", "Почему", "Пометка", "Оффер",
        "Запасной оффер", "Зацепка", "Кого искать", "Сайт", "ИНН", "Юрлицо", "Уверенность ИНН",
        "ОКВЭД", "Численность", "Доходы (ФНС)", "Регион", "Инфра-вакансий", "Висит с", "Вакансии",
        "Стек", "Зарплата от", "Зарплата до", "Ссылка на вакансию"] + LPR_COLS
WIDTH = {"Компания": 30, "Почему": 42, "Пометка": 34, "Оффер": 28, "Запасной оффер": 28,
         "Зацепка": 36, "Кого искать": 26, "Сайт": 24, "Юрлицо": 30, "ОКВЭД": 34, "Регион": 20,
         "Вакансии": 48, "Стек": 34, "Ссылка на вакансию": 30, "Исключение": 26,
         "Кандидаты ИНН": 60, "Комментарий": 30, "Следующий шаг": 20}
WRAP = {"Вакансии", "Почему", "Пометка", "Кандидаты ИНН"}

HOW = [
    ("Баллы (шаг 3)", ""),
    ("+3", "самая старая открытая инфра-вакансия висит 21+ день (от 2026-10-13)"),
    ("+3", "ищут лида или первого DevOps: head, lead, руководитель, тимлид, team lead, tech lead, "
           "архитектор, «единственный»"),
    ("+2", "2+ инфра-вакансии при численности до 1000"),
    ("+2", "численность 50–1000"),
    ("+1", "указана зарплата, верхняя граница ≥ 300 000"),
    ("+1", "в вакансиях или стеке есть явный триггер под оффер"),
    ("−3", "все инфра-вакансии junior / стажёр / trainee / ученик"),
    ("−2", "юрлицо не РФ: подтверждено на сайте или ТОО / LLC / Ltd / ЧК / Inc / B.V. в названии"),
    ("−2", "ОКВЭД 64.19 (банк) при численности > 1000"),
    ("", ""),
    ("Волны", ""),
    ("1", "балл ≥ 7"),
    ("2", "балл 4–6"),
    ("3", "балл ниже 4"),
    ("DevOps-фильтр", "в волны 1–2 только если хотя бы в одном названии вакансии есть devops, sre, "
                      "site reliability, platform engineer, платформ, kubernetes, k8s, devsecops, "
                      "mlops, ci/cd, релиз-инженер, «инфраструктуры и автоматизации»; иначе волна 3 "
                      "«сисадмин», а при Ansible / Terraform / Docker / Kubernetes в стеке — максимум 2"),
    ("Размер неизвестен", "размерные баллы = 0; прошла DevOps-фильтр и балл без размера ≥ 4 → "
                          "волна 2 «размер не проверен», иначе волна 3 «размер не определён». "
                          "Юрлицо с численностью < 10 по ФНС считается «размер неизвестен»"),
    ("Микро/малая", "численность 10–49 → волна 3"),
    ("Волна вручную", "если заполнена — главнее расчётной"),
    ("", ""),
    ("Офферы (шаг 4): главный — с наибольшим числом совпадений (название вакансии ×2, стек ×1)", ""),
] + [(str(n), f"{name}: " + ", ".join(p.pattern.replace("\\b", "").replace("\\w*", "…")
                                       .replace(".?", " ").split("|")))
     for n, name, p in OFFERS] + [
    ("MLOps", "если других триггеров нет — 9 или 10"),
    ("Без триггера", "10. Ускорение CI/CD, зацепка — первое название вакансии, +1 не начисляется"),
    ("", ""),
    ("Причины исключений", ""),
    ("гигант", "банки топ-20, телеком и «дочки», бигтех и ритейл, госкорпорации, сырьевые и "
               "промышленные холдинги, фастфуд; численность > 3000; «гигант (группа …)» — дочка "
               "крупной группы"),
    ("гос", "ГКУ, ГУП, ФГУП, ГАУ, ФКУ, ФБУН, министерства, НИИ, оборонные НПО, метрополитен, "
            "Гознак; ОКВЭД 84"),
    ("посредник", "ИП и ФИО, кадровые агентства, аутсорс/аутстафф, интеграторы; ОКВЭД 78"),
    ("посредник?", "интеграция, консалтинг, solutions, студии; ОКВЭД 70.22 — лист «Ручная проверка»"),
    ("инфра-вендор", "облака, хостинги, ЦОДы, anti-DDoS, ИБ-вендоры, ОС и виртуализация"),
    ("не DevOps", "все вакансии про поддержку, телефонию или 1С без DevOps-слов"),
    ("(решение вручную)", "исключено по решению от 14.10.2026"),
    ("", ""),
    ("Источники", "вакансии — hh.ru и getmatch; ИНН — DaData и сайты компаний; численность и "
                  "доходы — открытые данные ФНС (sshr, revexp) на 31.12.2025"),
]


def cell_value(r, col):
    if col == "ОКВЭД":
        return " ".join(x for x in (r.get("ОКВЭД код"), r.get("ОКВЭД расшифровка")) if x)
    if col == "Инфра-вакансий":
        return r.get("Вакансий инфраструктуры")
    if col == "Висит с":
        return as_date(r.get("Самая старая открытая")) or None
    if col == "Уверенность ИНН":
        return r.get("Уверенность сопоставления")
    if col == "Кандидаты ИНН":
        return r.get("Кандидаты (топ-3)")
    if col in ("Волна вручную", "Приоритет вручную"):
        return None
    if col in LPR_COLS and col != "Telegram ЛПР":
        return None
    v = r.get(col)
    return None if v in ("", None) else v


def write_sheet(wb, title, rows, cols, manual_col="Волна вручную",
                manual_opts='"1,2,3,исключить"'):
    from datetime import date
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    ws = wb.create_sheet(title)
    ws.append(cols)
    for r in rows:
        vals = [cell_value(r, c) for c in cols]
        ws.append([date.fromisoformat(v) if c == "Висит с" and v else v for c, v in zip(cols, vals)])
    head = PatternFill("solid", fgColor="DDE7F3")
    manual = PatternFill("solid", fgColor="FFF2CC")
    for i, c in enumerate(cols, 1):
        h = ws.cell(1, i)
        h.font, h.fill = Font(bold=True), head
        h.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(c, 14)
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell = row[0]
            cell.alignment = Alignment(wrap_text=c in WRAP, vertical="top")
            if c == "Висит с":
                cell.number_format = "DD.MM.YYYY"
            elif c in ("Доходы (ФНС)", "Зарплата от", "Зарплата до"):
                cell.number_format = "#,##0"
            elif c == manual_col:
                cell.fill = manual
            elif c in ("Сайт", "Ссылка на вакансию") and cell.value:
                cell.hyperlink = cell.value
            elif c in FLAGS:
                cell.alignment = Alignment(horizontal="center", vertical="top")
    if manual_col in cols:
        col = get_column_letter(cols.index(manual_col) + 1)
        dv = DataValidation(type="list", formula1=manual_opts, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(2, ws.max_row)}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def step_xlsx():
    from collections import Counter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    rows = json.loads((WORK / "step3.json").read_text(encoding="utf-8"))
    by_score = lambda r: (-(r.get("Балл") or 0), as_date(r.get("Самая старая открытая")) or "9")  # noqa: E731
    sheets = {k: [r for r in rows if r.get("Лист") == k]
              for k in ("Волна 1", "Волна 2", "Волна 3", "Исключены", "Ручная проверка")}
    for k in ("Волна 1", "Волна 2", "Ручная проверка"):
        sheets[k].sort(key=by_score)
    # волна 3: сверху «размер не определён» с баллом 3 — их пролистать глазами
    top3 = lambda r: "размер не определён" in (r.get("Пометка") or "") and r.get("Балл") == 3  # noqa: E731
    sheets["Волна 3"].sort(key=lambda r: (not top3(r),) + by_score(r))
    sheets["Исключены"].sort(key=lambda r: (r.get("Исключение") or "", r["Компания"]))

    wb = Workbook()
    wb.remove(wb.active)
    for k in ("Волна 1", "Волна 2", "Волна 3"):
        write_sheet(wb, k, sheets[k], COLS)
    write_sheet(wb, "Исключены", sheets["Исключены"],
                ["Компания", "Исключение"] + [c for c in COLS if c not in
                 ("Компания", "Волна", "Волна вручную", "Балл", "Почему", "Оффер", "Запасной оффер",
                  "Зацепка", "Кого искать") and c not in LPR_COLS])
    write_sheet(wb, "Ручная проверка", sheets["Ручная проверка"],
                COLS[:6] + ["Кандидаты ИНН"] + COLS[6:])
    ws = wb.create_sheet("Как считалось")
    for a, b in HOW:
        ws.append([a, b])
        if a and not b:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)

    print("листы:")
    for k, v in sheets.items():
        print(f"  {k:16s} {len(v)}")
    print("  волна 3 сверху (размер не определён, балл 3):", sum(top3(r) for r in sheets["Волна 3"]))
    print("исключены:", dict(Counter(r["Исключение"] for r in sheets["Исключены"]).most_common()))
    print(f"файл: {OUT}")


# ---------- v2: приоритет A/B/C по трём флагам (14.10.2026) ----------
OUT_V2 = EXPORTS / "b2b_v2.xlsx"
FLAGS = ["≤100 чел", "ОКВЭД разработка", "Доходы ≤1 млрд"]
COLS_V2 = (["Компания", "Приоритет"] + FLAGS +
           ["Приоритет вручную", "Численность", "ОКВЭД", "Доходы (ФНС)", "Сайт",
            "Ссылка на вакансию", "Зацепка", "Пометка", "ИНН", "Юрлицо", "Уверенность ИНН"]
           + LPR_COLS)
# пометки старой схемы волн — в v2 не нужны
OLD_NOTES = {"размер не проверен", "размер не определён", "микро/малая",
             "размер не определён, балл без размера ≥ 5"}
HOW_V2 = [
    ("Кто попадает", "все неисключённые компании (бывшие волны 1–3); исключения по гигантам, "
                     "госу, посредникам и инфра-вендорам — как раньше, лист «Исключены»"),
    ("", ""),
    ("Флаги (да / нет / ?)", ""),
    ("≤100 чел", "численность по открытым данным ФНС (sshr) на 31.12.2025 ≤ 100. Нет данных или "
                 "в юрлице < 10 человек → «?» (часто штат в другом юрлице группы)"),
    ("ОКВЭД разработка", "основной ОКВЭД 62.xx, 63.xx или 58.29; ИНН не найден → «?»"),
    ("Доходы ≤1 млрд", "доходы по открытым данным ФНС (revexp) за год ≤ 1 000 000 000 ₽; "
                       "нет данных → «?»"),
    ("", ""),
    ("Приоритет", ""),
    ("A", "все три флага «да»"),
    ("B", "два «да», третий «нет» или «?»"),
    ("C", "остальные"),
    ("Сортировка", "внутри приоритета — по численности по возрастанию, неизвестная численность в конце"),
    ("Приоритет вручную", "если заполнен — главнее расчётного (A / B / C / исключить)"),
    ("", ""),
    ("Ручная проверка", "«посредник?»: интеграция, консалтинг, solutions, студии, ОКВЭД 70.22"),
    ("", ""),
] + [x for x in HOW if x[0] in ("Причины исключений", "гигант", "гос", "посредник",
                                "инфра-вендор", "не DevOps", "(решение вручную)", "Источники")]


def flags_v2(r):
    staff = r.get("Численность")
    if staff is not None and staff < 10 or MANUAL.get(low(r["Компания"]), {}).get("size_unknown"):
        staff = None
    code, income = r.get("ОКВЭД код") or "", r.get("Доходы (ФНС)")
    f_staff = "?" if staff is None else "да" if staff <= 100 else "нет"
    f_okved = "?" if not code else "да" if re.match(r"6[23]\.|58\.29", code) else "нет"
    f_income = "?" if income is None else "да" if income <= 1e9 else "нет"
    yes = [f_staff, f_okved, f_income].count("да")
    return staff, (f_staff, f_okved, f_income), "A" if yes == 3 else "B" if yes == 2 else "C"


def old_manual(paths):
    """Ручные данные из прошлых выгрузок: {компания: {колонка: значение}}."""
    from openpyxl import load_workbook
    keep = ["Волна вручную", "Приоритет вручную"] + LPR_COLS
    out = {}
    for p in paths:
        if not p.exists():
            continue
        wb = load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            head = next(rows, None) or ()
            for row in rows:
                d = dict(zip(head, row))
                filled = {k: d[k] for k in keep if d.get(k) not in (None, "")}
                if d.get("Компания") and filled:
                    out.setdefault(d["Компания"], {}).update(filled)
    return out


def step_v2():
    from collections import Counter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    rows = json.loads((WORK / "step3.json").read_text(encoding="utf-8"))
    manual = old_manual([OUT, OUT_V2])
    for r in rows:
        r["_staff_v2"], (r[FLAGS[0]], r[FLAGS[1]], r[FLAGS[2]]), r["Приоритет"] = flags_v2(r)
        r["Пометка"] = "; ".join(n for n in (r.get("Пометка") or "").split("; ")
                                 if n and n not in OLD_NOTES)
        m = manual.get(r["Компания"], {})
        v = m.get("Приоритет вручную") or m.get("Волна вручную")
        # «Волна вручную» из v1 (1/2/3) не совпадает со шкалой A/B/C — переносим как есть с подписью
        r["_manual"] = (f"волна {v} (v1)" if v is not None and str(v) in ("1", "2", "3") else v)
        r["_lpr"] = {k: m[k] for k in LPR_COLS if k in m}
        if r.get("Лист") in ("Волна 1", "Волна 2", "Волна 3"):
            r["Лист v2"] = r["Приоритет"]
        else:
            r["Лист v2"] = r.get("Лист")
    order = lambda r: (r["_staff_v2"] is None, r["_staff_v2"] or 0, r["Компания"])  # noqa: E731
    sheets = {k: sorted((r for r in rows if r.get("Лист v2") == k), key=order)
              for k in ("A", "B", "C", "Ручная проверка")}
    sheets["Исключены"] = sorted((r for r in rows if r.get("Лист v2") == "Исключены"),
                                 key=lambda r: (r.get("Исключение") or "", r["Компания"]))

    wb = Workbook()
    wb.remove(wb.active)
    for k in ("A", "B", "C", "Ручная проверка", "Исключены"):
        cols = COLS_V2 if k != "Исключены" else ["Компания", "Исключение"] + COLS_V2[1:]
        ws = write_sheet(wb, k, sheets[k], cols, manual_col="Приоритет вручную",
                         manual_opts='"A,B,C,исключить"')
        for i, r in enumerate(sheets[k], 2):    # ручные данные — поверх пустых ячеек
            for j, c in enumerate(cols, 1):
                if c == "Приоритет вручную" and r["_manual"] is not None:
                    ws.cell(i, j, r["_manual"])
                elif c in r["_lpr"]:
                    ws.cell(i, j, r["_lpr"][c])
    ws = wb.create_sheet("Как считалось")
    for a, b in HOW_V2:
        ws.append([a, b])
        if a and not b:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT_V2)

    print("листы:")
    for k, v in sheets.items():
        print(f"  {k:16s} {len(v)}")
    for k in ("A", "B", "C"):
        print(f"  {k}: флаги", dict(Counter(tuple(r[f] for f in FLAGS) for r in sheets[k]).most_common(6)))
    print("перенесено ручных записей:", sum(1 for r in rows if r["_manual"] is not None or r["_lpr"]))
    print("\nпервые 30 из A:")
    for r in sheets["A"][:30]:
        inc = f"{r['Доходы (ФНС)'] / 1e6:,.0f} млн".replace(",", " ")
        print(f"  {r['Компания'][:34]:34s} | {r['Численность']:>4} | "
              f"{r['ОКВЭД код']:8s} {(r.get('ОКВЭД расшифровка') or '')[:38]:38s} | {inc:>9s}")
    print(f"файл: {OUT_V2}")


# ---------- v3: ядро — до 500 человек, доходы до 3 млрд, ОКВЭД из списка (14.10.2026) ----------
OUT_V3 = EXPORTS / "b2b_v3.xlsx"
# код (префикс) → (группа, кто там, офферы); проверяются по порядку, первый подходящий
OKVED_V3 = [
    ("62.01", "ИТ", "разработка ПО: продуктовые компании, SaaS, студии", [7, 8, 9, 10, 2]),
    ("62.0", "ИТ", "разработка ПО: продуктовые компании, SaaS, студии", [7, 8, 9, 10, 2]),
    ("58.29", "ИТ", "издатели ПО, вендоры коробочных и облачных продуктов", [10, 9, 3]),
    ("58.21", "ИТ", "геймдев", [4, 2, 1]),
    ("63.11", "ИТ", "обработка данных, хостинг, SaaS, облачные сервисы", [1, 2, 4, 6]),
    ("63.12", "ИТ", "веб-порталы, маркетплейсы, агрегаторы, сервисы объявлений", [1, 2, 4]),
    ("62.09", "ИТ", "прочая ИТ-деятельность, часто стартапы с размытым профилем", [7, 10]),
    ("62.03", "ИТ", "управление компьютерным оборудованием", [6, 5]),
    ("47.91", "нагрузка", "интернет-магазины, e-com (сезонные пики)", [1, 2, 4]),
    ("64.92", "нагрузка", "онлайн-МФО", [3, 1, 5]),
    ("66.19", "нагрузка", "платёжные сервисы, процессинг, финтех", [3, 1, 2]),
    ("66.12", "нагрузка", "брокеры, инвест-платформы", [3, 1]),
    ("61.10", "нагрузка", "провайдеры, телеком, VPN, CDN", [6, 2, 5]),
    ("61.90", "нагрузка", "провайдеры, телеком, VPN, CDN", [6, 2, 5]),
    ("73.11", "нагрузка", "adtech, рекламные платформы, аналитика", [4, 2, 9]),
    ("73.20", "нагрузка", "adtech, рекламные платформы, аналитика", [4, 2, 9]),
    ("85.41", "нагрузка", "EdTech, онлайн-школы", [4, 7, 2]),
    ("86.90", "нагрузка", "телемедицина, медтех (152-ФЗ, переезд в РФ-облака)", [3, 5]),
    ("79.11", "нагрузка", "онлайн-бронирование, тревел-сервисы", [1, 4]),
    ("79.90", "нагрузка", "онлайн-бронирование, тревел-сервисы", [1, 4]),
    ("53.20", "нагрузка", "доставка, такси-агрегаторы, логистика с приложением", [1, 2, 9]),
    ("49.32", "нагрузка", "доставка, такси-агрегаторы, логистика с приложением", [1, 2, 9]),
]
OFFER_NAMES = {n: name for n, name, _ in OFFERS}
FLAGS_V3 = ["≤500 чел", "ОКВЭД из списка", "Доходы ≤3 млрд"]
COLS_V3 = (["Компания", "Приоритет"] + FLAGS_V3 +
           ["Приоритет вручную", "Сегмент", "Офферы по ОКВЭД", "Численность", "ОКВЭД",
            "Доходы (ФНС)", "Сайт", "Ссылка на вакансию", "Зацепка", "Пометка", "ИНН", "Юрлицо",
            "Уверенность ИНН"] + LPR_COLS)


def okved_v3(code):
    """(группа, кто там, офферы) для основного ОКВЭД или None. «62.0» у DaData — это 62.01."""
    code = code or ""
    for pref, group, who, offers in OKVED_V3:
        if code == pref or code.startswith(pref + ".") or (pref == "62.0" and code == "62.0") \
                or (len(pref) == 5 and code.startswith(pref)):
            return group, who, offers
    return None


def flags_v3(r):
    staff = r.get("Численность")
    if staff is not None and staff < 10 or MANUAL.get(low(r["Компания"]), {}).get("size_unknown"):
        staff = None
    code, income = r.get("ОКВЭД код") or "", r.get("Доходы (ФНС)")
    seg = okved_v3(code)
    f_staff = "?" if staff is None else "да" if staff <= 500 else "нет"
    f_okved = "?" if not code else "да" if seg else "нет"
    f_income = "?" if income is None else "да" if income <= 3e9 else "нет"
    yes = [f_staff, f_okved, f_income].count("да")
    return staff, seg, (f_staff, f_okved, f_income), "A" if yes == 3 else "B" if yes == 2 else "C"


def step_v3():
    from collections import Counter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    rows = json.loads((WORK / "step3.json").read_text(encoding="utf-8"))
    manual = old_manual([OUT, OUT_V2, OUT_V3])
    for r in rows:
        r["_staff_v2"], seg, flags, r["Приоритет"] = flags_v3(r)
        r[FLAGS_V3[0]], r[FLAGS_V3[1]], r[FLAGS_V3[2]] = flags
        r["Сегмент"] = f"{seg[0]}: {seg[1]}" if seg else ""
        r["Офферы по ОКВЭД"] = "\n".join(f"{n}. {OFFER_NAMES[n]}" for n in seg[2]) if seg else ""
        r["Пометка"] = "; ".join(n for n in (r.get("Пометка") or "").split("; ")
                                 if n and n not in OLD_NOTES)
        m = manual.get(r["Компания"], {})
        v = m.get("Приоритет вручную") or m.get("Волна вручную")
        r["_manual"] = (f"волна {v} (v1)" if v is not None and str(v) in ("1", "2", "3") else v)
        r["_lpr"] = {k: m[k] for k in LPR_COLS if k in m}
        r["Лист v2"] = (r["Приоритет"] if r.get("Лист") in ("Волна 1", "Волна 2", "Волна 3")
                        else r.get("Лист"))
    order = lambda r: (r["_staff_v2"] is None, r["_staff_v2"] or 0, r["Компания"])  # noqa: E731
    sheets = {k: sorted((r for r in rows if r.get("Лист v2") == k), key=order)
              for k in ("A", "B", "C", "Ручная проверка")}
    sheets["Исключены"] = sorted((r for r in rows if r.get("Лист v2") == "Исключены"),
                                 key=lambda r: (r.get("Исключение") or "", r["Компания"]))

    WIDTH.update({"Сегмент": 34, "Офферы по ОКВЭД": 36})
    WRAP.add("Офферы по ОКВЭД")
    FLAGS.extend(f for f in FLAGS_V3 if f not in FLAGS)    # центрирование флагов в write_sheet
    wb = Workbook()
    wb.remove(wb.active)
    for k in ("A", "B", "C", "Ручная проверка", "Исключены"):
        cols = COLS_V3 if k != "Исключены" else ["Компания", "Исключение"] + COLS_V3[1:]
        ws = write_sheet(wb, k, sheets[k], cols, manual_col="Приоритет вручную",
                         manual_opts='"A,B,C,исключить"')
        for i, r in enumerate(sheets[k], 2):
            for j, c in enumerate(cols, 1):
                if c == "Приоритет вручную" and r["_manual"] is not None:
                    ws.cell(i, j, r["_manual"])
                elif c in r["_lpr"]:
                    ws.cell(i, j, r["_lpr"][c])
    how = [
        ("Кто попадает", HOW_V2[0][1]),
        ("", ""),
        ("Флаги (да / нет / ?)", ""),
        ("≤500 чел", "численность по ФНС (sshr) на 31.12.2025 ≤ 500. Нет данных или в юрлице "
                     "< 10 человек → «?»"),
        ("ОКВЭД из списка", "основной ОКВЭД есть в таблице ниже (ИТ-ядро или высоконагруженные "
                            "не-ИТ); ИНН не найден → «?»"),
        ("Доходы ≤3 млрд", "доходы по ФНС (revexp) ≤ 3 000 000 000 ₽: деньги есть, своей "
                           "платформенной команды ещё нет; нет данных → «?»"),
        ("", ""),
        ("Приоритет", ""),
        ("A", "все три флага «да»"),
        ("B", "два «да», третий «нет» или «?»"),
        ("C", "остальные"),
        ("Сортировка", "внутри приоритета — по численности по возрастанию, неизвестная в конце"),
        ("Приоритет вручную", "если заполнен — главнее расчётного (A / B / C / исключить)"),
        ("", ""),
        ("ОКВЭД → сегмент → офферы", ""),
    ] + [(pref, f"[{group}] {who} → офферы {', '.join(map(str, offers))}")
         for pref, group, who, offers in OKVED_V3 if pref != "62.0"] + [
        ("", ""),
        ("Офферы", ""),
    ] + [(str(n), name) for n, name, _ in OFFERS] + [("", "")] + HOW_V2[-9:]
    ws = wb.create_sheet("Как считалось")
    for a, b in how:
        ws.append([a, b])
        if a and not b:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT_V3)

    print("листы:")
    for k, v in sheets.items():
        print(f"  {k:16s} {len(v)}")
    for k in ("A", "B", "C"):
        print(f"  {k}: флаги", dict(Counter(tuple(r[f] for f in FLAGS_V3)
                                            for r in sheets[k]).most_common(6)))
    print("  A по сегментам:", dict(Counter(r["Сегмент"].split(":")[0] for r in sheets["A"])))
    print("  A по ОКВЭД:", dict(Counter((r.get("ОКВЭД код") or "")[:5] for r in sheets["A"]).most_common(10)))
    print("перенесено ручных записей:", sum(1 for r in rows if r["_manual"] is not None or r["_lpr"]))
    print("\nпервые 30 из A:")
    for r in sheets["A"][:30]:
        inc = f"{r['Доходы (ФНС)'] / 1e6:,.0f} млн".replace(",", " ")
        print(f"  {r['Компания'][:32]:32s} | {r['Численность']:>4} | {r['ОКВЭД код']:8s} | {inc:>9s} | "
              + ",".join(o.split(".")[0] for o in r["Офферы по ОКВЭД"].split("\n")))
    print(f"файл: {OUT_V3}")


if __name__ == "__main__":
    a = sys.argv[1:]
    if a[:1] == ["step1"]:
        step1()
    elif a[:1] == ["step2"]:
        step2(int(a[a.index("--limit") + 1]) if "--limit" in a else None)
    elif a[:1] == ["step2b"]:
        step2b()
    elif a[:1] == ["step3"]:
        step3()
    elif a[:1] == ["xlsx"]:
        step_xlsx()
    elif a[:1] == ["v2"]:
        step_v2()
    elif a[:1] == ["v3"]:
        step_v3()
    else:
        print(__doc__)
