"""Таблица компаний для B2B-продаж DevOps-услуг, xlsx.

    python scripts/b2b_xlsx.py

Одна строка — одна компания, которая сейчас ищет инфраструктурных
инженеров: DevOps, SRE, DevSecOps, сисадмин, сетевой, платформа,
инфраструктура. Считаются только такие вакансии и только живые (без
«В архиве с…»): по ним же даты, вилка, стек и ссылка. Компании без
инфра-вакансий, кадровые агентства, стоп-лист аккаунта и «Компания скрыта»
в таблицу не попадают. Подразделения одной компании склеиваются: «Сбер» и
«Сбер. IT», «МТС» и «"МТС", Работа в IT», «Cloud.ru: Инженерия» и
«Cloud.ru: Разработка».

Приоритет: A — 3+ инфра-вакансии и последняя не старше 30 дней; B — 2+
вакансии; C — одна вакансия.

Пишет exports/b2b_<дата>.xlsx в папке первого аккаунта: лист «A+B» и лист
«Все», шапка закреплена, фильтры включены. Базы только читает.
"""
import collections
import datetime as dt
import re
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path("/Users/mr_macbook/Documents")
ACCS = ("hh_autoreply", "hh_autoreply_2", "hh_autoreply_3")
OUT_DIR = BASE / "hh_autoreply" / "exports"
TODAY = dt.date.today()

INFRA_RE = re.compile(
    r"devops|dev ops|devsecops|\bsre\b|site reliab|сисадмин|системн\w*\s+администратор"
    r"|system\s*admin|sysadmin|сетев|network|платформ|platform|инфраструктур"
    r"|infrastructure", re.I)
ARCHIVE_RE = re.compile(r"в архиве с", re.I)
AGENCY_RE = re.compile(r"кадров|рекрут|recruit|\bhr\b|агентств|staff|персонал|"
                       r"headhunt|getmatch agency|подбор", re.I)
HIDDEN_RE = re.compile(r"^компания скрыта", re.I)
# хвосты подразделений: «Сбер. IT», «"МТС", Работа в IT», «Cloud.ru: Инженерия»
DIVISION_RE = re.compile(
    r"\s*[.,:]\s*(it|ит|ит-специалисты|работа в it|it & digital|digital|разработка|"
    r"инженерия|операции|технологии.*|направление.*|data science|it-специалисты|"
    r"технический блок|центральный офис)\s*$", re.I)
FORM_RE = re.compile(r"^(ооо|оао|зао|пао|ао|ип|гку|гуп|нао|ткп|фгуп|гбу)\s+", re.I)
LAT2CYR = str.maketrans("aeopcxykmthbAEOPCXYKMTHB", "аеорсхукмтнвАЕОРСХУКМТНВ")
# не стек, а обязанности или куски текста
SKILL_STOP_RE = re.compile(
    r":|чем предстоит|robotics|настройка пк|ремонт пк|сборка пк|техническая поддержка"
    r"|установка по|настройка по\b|обучение|работа в команде|^devops$|^it$", re.I)


def display_name(name):
    n = re.sub(r"\s+", " ", (name or "").replace("\xa0", " ")).strip()
    n = DIVISION_RE.sub("", n).strip(" ,.:")
    return re.sub(r'^"([^"]+)"$', r"\1", n)          # «"МТС"» → «МТС»


def key_of(name):
    """Ключ склейки. В словах, где кириллица и латиница вперемешку («Aльфа»),
    латиница переводится в кириллицу; чисто латинские названия не трогаем."""
    k = display_name(name).lower()
    k = FORM_RE.sub("", k)
    k = re.sub(r"[«»\"'`]", "", k)
    words = []
    for w in re.split(r"[^\w]+", k):
        if re.search(r"[а-я]", w) and re.search(r"[a-z]", w):
            w = w.translate(LAT2CYR)
        if w:
            words.append(w)
    return " ".join(words)


def blacklist():
    try:
        lines = (BASE / "hh_autoreply" / "blacklist.txt").read_text(encoding="utf-8").splitlines()
    except OSError:
        return []
    return [l.strip().lower() for l in lines if l.strip() and not l.startswith("#")]


def as_date(s):
    try:
        return dt.date.fromisoformat((s or "")[:10])
    except ValueError:
        return None


def clean_title(t):
    return re.sub(r"\s+", " ", (t or "").split("\n")[0]).strip()


def collect():
    comp = collections.defaultdict(lambda: {"names": collections.Counter(), "vac": {}})
    for acc in ACCS:
        db = sqlite3.connect(BASE / acc / "hh_responses.db")
        facts = {r[0]: r[1:] for r in db.execute(
            "SELECT vacancy_id, salary_from, salary_to, currency, published_at FROM vacancy_facts")}
        skills = collections.defaultdict(list)
        for vid, skill, st in db.execute("SELECT vacancy_id, skill, status FROM vacancy_skills"):
            if st != "stale" and skill:
                skills[vid].append(skill)
        for vid, url, title, company, ts in db.execute(
                "SELECT id, url, title, company, ts FROM responses"):
            if not company or not title or ARCHIVE_RE.search(title):
                continue
            t = clean_title(title)
            if not INFRA_RE.search(t):
                continue
            f = facts.get(vid) or (None, None, None, None)
            rub = (f[2] or "RUR").upper() in ("RUR", "RUB", "₽")
            c = comp[key_of(company)]
            c["names"][display_name(company)] += 1
            c["vac"].setdefault(f"hh{vid}", {
                "title": t, "url": url, "date": as_date(f[3]) or as_date(ts),
                "from": f[0] if rub else None, "to": f[1] if rub else None,
                "skills": skills.get(vid, [])})
        if "gm_vacancies" in {r[0] for r in db.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}:
            for vid, url, pos, company, pub, s_from, s_to, cur, sk in db.execute(
                    "SELECT id, url, position, company, published_at, salary_from, salary_to, "
                    "salary_currency, skills FROM gm_vacancies"):
                t = clean_title(pos)
                if not company or not INFRA_RE.search(t):
                    continue
                rub = (cur or "").upper() in ("RUR", "RUB", "₽")
                c = comp[key_of(company)]
                c["names"][display_name(company)] += 1
                c["vac"].setdefault(f"gm{vid}", {
                    "title": t, "url": url, "date": as_date(pub),
                    "from": s_from if rub else None, "to": s_to if rub else None,
                    "skills": [s for s in (sk or "").split(",") if s]})
    return comp


def build_rows(comp):
    stop = blacklist()
    rows = []
    for key, c in comp.items():
        if not key or not c["vac"]:
            continue
        name = c["names"].most_common(1)[0][0]
        low = name.lower()
        if HIDDEN_RE.match(name) or AGENCY_RE.search(name) or any(s in low for s in stop):
            continue
        vac = sorted(c["vac"].values(), key=lambda v: v["date"] or dt.date.min, reverse=True)
        dates = [v["date"] for v in vac if v["date"]]
        last = max(dates) if dates else None
        oldest = min(dates) if dates else None
        titles = list(dict.fromkeys(v["title"] for v in vac))[:5]
        stack = collections.Counter()
        for v in vac:
            seen = set()
            for s in v["skills"]:
                s = s.strip()
                if not s or len(s) > 25 or SKILL_STOP_RE.search(s) or s.lower() in seen:
                    continue
                # технологии почти всегда латиницей; русская фраза из нескольких
                # слов — это обязанность («Настройка роутеров»), а не стек
                if not re.search(r"[A-Za-z]", s) and " " in s:
                    continue
                seen.add(s.lower())
                stack[s] += 1
        froms = [v["from"] for v in vac if v["from"] and 50_000 <= v["from"] <= 1_500_000]
        tos = [v["to"] for v in vac if v["to"] and 50_000 <= v["to"] <= 1_500_000]
        fresh = last is not None and (TODAY - last).days <= 30
        n = len(vac)
        prio = "A" if n >= 3 and fresh else "B" if n >= 2 else "C"
        rows.append({
            "Компания": name,
            "Приоритет": prio,
            "Вакансий инфраструктуры": n,
            "Вакансии (до 5)": "\n".join(titles),
            "Последняя инфра-вакансия": last,
            "Самая старая открытая": oldest,
            "Стек": ", ".join(s for s, _ in stack.most_common(8)),
            "Зарплата от": min(froms) if froms else None,
            "Зарплата до": max(tos) if tos else None,
            "Telegram ЛПР": "",
            "Сайт": "",
            "Ссылка на вакансию": next((v["url"] for v in vac if v["url"]), ""),
        })
    rows.sort(key=lambda r: (r["Приоритет"], -r["Вакансий инфраструктуры"],
                             -(r["Последняя инфра-вакансия"] or dt.date.min).toordinal()))
    return rows


WIDTH = {"Компания": 32, "Приоритет": 10, "Вакансий инфраструктуры": 13, "Вакансии (до 5)": 55,
         "Последняя инфра-вакансия": 14, "Самая старая открытая": 14, "Стек": 45,
         "Зарплата от": 12, "Зарплата до": 12, "Telegram ЛПР": 18, "Сайт": 22,
         "Ссылка на вакансию": 45}
FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "EDEDED"}


def write_sheet(ws, rows):
    cols = list(WIDTH)
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = WIDTH[c]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        prio = row[cols.index("Приоритет")]
        prio.fill = PatternFill("solid", fgColor=FILL.get(prio.value, "FFFFFF"))
        for name in ("Последняя инфра-вакансия", "Самая старая открытая"):
            row[cols.index(name)].number_format = "DD.MM.YYYY"
        for name in ("Зарплата от", "Зарплата до"):
            row[cols.index(name)].number_format = "# ##0"
        link = row[cols.index("Ссылка на вакансию")]
        if link.value:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def main():
    rows = build_rows(collect())
    wb = Workbook()
    ws = wb.active
    ws.title = "A+B"
    write_sheet(ws, [r for r in rows if r["Приоритет"] in ("A", "B")])
    write_sheet(wb.create_sheet("Все"), rows)
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f"b2b_{TODAY.isoformat()}.xlsx"
    wb.save(out)
    cnt = collections.Counter(r["Приоритет"] for r in rows)
    print(f"компаний: {len(rows)} (A {cnt['A']}, B {cnt['B']}, C {cnt['C']})")
    print(f"файл: {out}")
    for r in rows[:15]:
        print(f"  {r['Приоритет']} {r['Вакансий инфраструктуры']:3d}  {r['Компания'][:34]:34s} "
              f"{r['Последняя инфра-вакансия']}  {(r['Стек'] or '')[:40]}")


if __name__ == "__main__":
    main()
